# -*- coding: utf-8 -*-
"""Importador de documentos Word (.docx) a paginas de la Wiki.

Convierte parrafos/listas/tablas/encabezados de un .docx al mismo formato de
bloques que usa el editor de wiki (EditorJS). Negrita/cursiva/subrayado/
tachado/enlaces se preservan como HTML embebido en el texto de cada bloque --
wiki_store._clean() lo sanea con la misma lista blanca que usa el contenido
escrito a mano en el editor, asi que es seguro generarlo aqui sin escapar a
ciegas. Las imagenes incrustadas no se importan (el editor de wiki todavia no
tiene una herramienta de imagen) -- se cuentan y se devuelven como aviso."""
from __future__ import annotations

import re
from html import escape

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

IMPORT_DOCX_MAX_SIZE_MB = 15

_W_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
_HEADING_RE = re.compile(r"^Heading (\d+)$", re.I)


def _iter_block_items(doc: Document):
    """Recorre parrafos y tablas en el orden real del documento (python-docx
    los expone por separado via doc.paragraphs/doc.tables, perdiendo el
    orden y el entrelazado entre ambos)."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _run_html(run) -> str:
    text = run.text
    if not text:
        return ""
    html = escape(text)
    if run.bold:
        html = f"<strong>{html}</strong>"
    if run.italic:
        html = f"<em>{html}</em>"
    if run.underline:
        html = f"<u>{html}</u>"
    if run.font.strike:
        html = f"<s>{html}</s>"
    return html


def _hyperlink_html(link: Hyperlink) -> str:
    inner = "".join(_run_html(r) for r in link.runs) or escape(link.text or "")
    url = (link.address or "").strip()
    if not url:
        return inner
    return f'<a href="{escape(url)}">{inner}</a>'


def _paragraph_html(paragraph: Paragraph) -> str:
    parts = []
    for item in paragraph.iter_inner_content():
        parts.append(_hyperlink_html(item) if isinstance(item, Hyperlink) else _run_html(item))
    return "".join(parts)


def _heading_level(style_name: str) -> int | None:
    """Nivel 1-4 (el header tool del editor solo admite ese rango) o None si
    el estilo no es un encabezado."""
    if not style_name:
        return None
    name = style_name.strip()
    if name.lower() == "title":
        return 1
    m = _HEADING_RE.match(name)
    return min(int(m.group(1)), 4) if m else None


def _num_fmt(doc: Document, num_id: str) -> str | None:
    """Resuelve el formato (bullet/decimal/...) de una lista aplicada por
    numeracion directa (numPr) en vez de un estilo con nombre -- caso comun
    cuando la lista se creo con los botones de la cinta de Word en vez del
    galeria de estilos. Cualquier fallo se trata como 'no se pudo resolver'
    (se usara vineta por defecto) en vez de romper la importacion."""
    try:
        numbering = doc.part.numbering_part.element
        num = numbering.find(f'w:num[@w:numId="{num_id}"]', _W_NS)
        if num is None:
            return None
        ref = num.find("w:abstractNumId", _W_NS)
        if ref is None:
            return None
        abstract_id = ref.get(f'{{{_W_NS["w"]}}}val')
        abstract = numbering.find(f'w:abstractNum[@w:abstractNumId="{abstract_id}"]', _W_NS)
        if abstract is None:
            return None
        lvl0 = abstract.find('w:lvl[@w:ilvl="0"]', _W_NS)
        if lvl0 is None:
            return None
        fmt = lvl0.find("w:numFmt", _W_NS)
        return fmt.get(f'{{{_W_NS["w"]}}}val') if fmt is not None else None
    except Exception:
        return None


def _list_kind(paragraph: Paragraph, doc: Document) -> str | None:
    """'ordered' | 'unordered' | None (el parrafo no es un item de lista)."""
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("List Number"):
        return "ordered"
    if style_name.startswith("List Bullet"):
        return "unordered"
    pPr = paragraph._p.pPr
    num_pr = pPr.numPr if pPr is not None else None
    if num_pr is None or num_pr.numId is None:
        return None
    fmt = _num_fmt(doc, str(num_pr.numId.val))
    return "unordered" if fmt in (None, "bullet") else "ordered"


def _table_block(table: Table) -> dict:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cells.append("<br>".join(
                _paragraph_html(p) for p in cell.paragraphs if p.text.strip()
            ))
        rows.append(cells)
    return {"type": "table", "data": {"withHeadings": len(rows) > 1, "content": rows}}


def import_docx_to_editorjs(fileobj) -> tuple[dict, str, list[str]]:
    """Lee un .docx y devuelve (content_editorjs, titulo_sugerido, avisos).
    No falla ante contenido inusual -- lo reporta como aviso y sigue, igual
    que import_tasks_from_excel."""
    nombre_archivo = getattr(fileobj, "name", "") or ""
    if not nombre_archivo.lower().endswith(".docx"):
        raise ValueError("el archivo debe ser un .docx")
    tamano = getattr(fileobj, "size", 0) or 0
    if tamano > IMPORT_DOCX_MAX_SIZE_MB * 1024 * 1024:
        raise ValueError(f"el archivo pesa demasiado (maximo {IMPORT_DOCX_MAX_SIZE_MB} MB)")

    try:
        doc = Document(fileobj)
    except Exception as e:
        raise ValueError(f"no se pudo leer el documento (¿es un .docx valido?): {e}")

    blocks: list[dict] = []
    title = ""
    list_buffer: list[str] = []
    list_kind: str | None = None

    def flush_list():
        nonlocal list_buffer, list_kind
        if list_buffer:
            blocks.append({"type": "list", "data": {"style": list_kind or "unordered", "items": list_buffer}})
        list_buffer, list_kind = [], None

    for item in _iter_block_items(doc):
        if isinstance(item, Table):
            flush_list()
            if any(cell.text.strip() for row in item.rows for cell in row.cells):
                blocks.append(_table_block(item))
            continue

        text = item.text.strip()
        style_name = item.style.name if item.style else ""

        lvl = _heading_level(style_name)
        if lvl:
            flush_list()
            if text:
                if not title:
                    title = text
                blocks.append({"type": "header", "data": {"text": _paragraph_html(item), "level": lvl}})
            continue

        kind = _list_kind(item, doc)
        if kind:
            if not text:
                continue
            if list_buffer and list_kind != kind:
                flush_list()
            list_kind = kind
            list_buffer.append(_paragraph_html(item))
            continue

        flush_list()
        if not text:
            continue
        if style_name.strip().lower() in ("quote", "intense quote"):
            blocks.append({"type": "quote", "data": {"text": _paragraph_html(item), "caption": "", "alignment": "left"}})
        else:
            blocks.append({"type": "paragraph", "data": {"text": _paragraph_html(item)}})

    flush_list()

    avisos: list[str] = []
    n_imgs = len(doc.inline_shapes)
    if n_imgs:
        avisos.append(
            f"se omitieron {n_imgs} imagen{'es' if n_imgs != 1 else ''} incrustada"
            f"{'s' if n_imgs != 1 else ''} (el editor de wiki todavia no soporta imagenes)"
        )
    if not blocks:
        avisos.append("el documento no tenia contenido reconocible, la pagina se creo vacia")

    if not title:
        title = nombre_archivo.rsplit(".", 1)[0].strip() or "Documento importado"

    content = {"time": 0, "blocks": blocks, "version": "2.29.0"}
    return content, title, avisos
