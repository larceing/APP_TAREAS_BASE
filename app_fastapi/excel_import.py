# -*- coding: utf-8 -*-
"""Importador de tareas desde Excel (.xlsx) para el Backlog. Copiado de
core/views.py::import_tasks_from_excel -- puro (sin Django, sin dataclases
de Task: aqui trabaja sobre listas de tuplas (id, nombre) y devuelve dicts
listos para pasarle a db.crear_tarea). No falla ante filas/columnas
invalidas -- las reporta como avisos y sigue con el resto de filas."""
from __future__ import annotations

import unicodedata
from datetime import datetime


IMPORT_EXCEL_MAX_SIZE_MB = 2
IMPORT_EXCEL_MAX_ROWS = 500


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def _norm(s) -> str:
    s = "" if s is None else str(s).strip()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return s.lower()


def _match_by_name(valor: str, opciones: list):
    """opciones: lista de (id_o_nombre, nombre). Devuelve (id_o_nombre, encontrado)."""
    if not valor:
        return "", True
    objetivo = _norm(valor)
    for id_, nombre in opciones:
        if _norm(nombre) == objetivo:
            return id_, True
    return "", False


def _parse_excel_fecha(valor) -> str:
    if valor is None or valor == "":
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d")
    s = str(valor).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def import_tasks_from_excel(
    fileobj, nombre_archivo: str, tamano: int, titulos_existentes: set,
    categorias: list, sistemas: list, empresas: list, iniciativas: list, priorities: list,
) -> tuple:
    """Lee un .xlsx (fila 1 = cabeceras Titulo/Descripcion/Categoria/Sistema/
    Empresa/Iniciativa/Prioridad/Fecha Inicio/Fecha Fin/Proximo paso) y
    devuelve (nuevas, avisos), donde 'nuevas' es una lista de dicts con las
    claves que espera db.crear_tarea (category/title/description/next_step/
    priority/fecha_inicio/fecha_fin/empresa_id/sistema_id/iniciativa_id).
    titulos_existentes: set ya normalizado (via _norm) de titulos ya usados
    en el proyecto, para no duplicar."""
    if not nombre_archivo.lower().endswith(".xlsx"):
        raise ValueError("el archivo debe ser un .xlsx (usa la plantilla descargable)")
    if tamano > IMPORT_EXCEL_MAX_SIZE_MB * 1024 * 1024:
        raise ValueError(f"el archivo pesa demasiado (maximo {IMPORT_EXCEL_MAX_SIZE_MB} MB)")

    import openpyxl  # import diferido: solo se carga al usarse, no en el arranque de la app

    wb = openpyxl.load_workbook(fileobj, data_only=True)
    ws = wb.active

    header_row = [(_norm(c.value), i) for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
    col_idx = {name: i for name, i in header_row}
    if _norm("Titulo") not in col_idx:
        raise ValueError("no se encontro la columna 'Titulo' -- ¿es la plantilla correcta?")

    n_filas = max(ws.max_row - 1, 0)
    if n_filas > IMPORT_EXCEL_MAX_ROWS:
        raise ValueError(f"el archivo tiene {n_filas} filas, el maximo por importacion es {IMPORT_EXCEL_MAX_ROWS}")

    def get(row, col_name):
        idx = col_idx.get(_norm(col_name))
        return row[idx].value if idx is not None and idx < len(row) else None

    cat_opciones = [(c, c) for c in categorias]
    sis_opciones = [(s["id"], s["nombre"]) for s in sistemas]
    emp_opciones = [(e["empresa_id"], e["nombre"]) for e in empresas]
    ini_opciones = [(i["id"], i["nombre"]) for i in iniciativas]

    titulos_vistos = set(titulos_existentes)
    nuevas = []
    avisos = []

    for n, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if all(c.value in (None, "") for c in row):
            continue
        titulo = _clean(get(row, "Titulo"))
        if not titulo:
            avisos.append(f"fila {n}: sin Titulo, omitida")
            continue
        if _norm(titulo) in titulos_vistos:
            avisos.append(f"fila {n} ({titulo}): ya existe una tarea con este titulo en el proyecto, omitida")
            continue

        categoria, cat_ok = _match_by_name(_clean(get(row, "Categoria")), cat_opciones)
        if not cat_ok:
            avisos.append(f"fila {n} ({titulo}): Categoria '{get(row, 'Categoria')}' no encontrada, se deja vacia")
        sistema_id, sis_ok = _match_by_name(_clean(get(row, "Sistema")), sis_opciones)
        if not sis_ok:
            avisos.append(f"fila {n} ({titulo}): Sistema '{get(row, 'Sistema')}' no encontrado, se deja vacio")
        empresa_id, emp_ok = _match_by_name(_clean(get(row, "Empresa")), emp_opciones)
        if not emp_ok:
            avisos.append(f"fila {n} ({titulo}): Empresa '{get(row, 'Empresa')}' no encontrada, se deja vacia")
        iniciativa_id, ini_ok = _match_by_name(_clean(get(row, "Iniciativa")), ini_opciones)
        if not ini_ok:
            avisos.append(f"fila {n} ({titulo}): Iniciativa '{get(row, 'Iniciativa')}' no encontrada, se deja vacia")

        prioridad_raw = _clean(get(row, "Prioridad"))
        prioridad = next((p for p in priorities if _norm(p) == _norm(prioridad_raw)), None)
        if prioridad_raw and prioridad is None:
            avisos.append(f"fila {n} ({titulo}): Prioridad '{prioridad_raw}' no valida, se usa Normal")
        prioridad = prioridad or "Normal"

        nuevas.append({
            "title": titulo, "category": categoria,
            "description": _clean(get(row, "Descripcion")), "next_step": _clean(get(row, "Proximo paso")),
            "priority": prioridad,
            "fecha_inicio": _parse_excel_fecha(get(row, "Fecha Inicio")),
            "fecha_fin": _parse_excel_fecha(get(row, "Fecha Fin")),
            "empresa_id": empresa_id, "sistema_id": sistema_id, "iniciativa_id": iniciativa_id,
        })
        titulos_vistos.add(_norm(titulo))

    return nuevas, avisos
