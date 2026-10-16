# -*- coding: utf-8 -*-
"""Pagina de Mantis, copiada de core/views.py::mantis_page/mantis_incidencias/
mantis_importar_tareas + core/templates/mantis.html.

Ojo con el patron de carga: la pagina en si (GET /mantis/) NO llama a Mantis
-- solo pinta el hueco vacio ("Cargando...") y el JS de la plantilla pide
las incidencias aparte, de forma asincrona, a /mantis/incidencias/json. El
pool "Incidencias" tiene mucho historico y tarda bastante en paginarse
entero (~15-20s medido contra el real), asi que si esa llamada fuera
sincrona DENTRO de la ruta (como estaba antes en esta v2), toda la pagina
se quedaria en blanco ese rato -- exactamente el bug que hizo que este
"Mantis" pareciera mas lento y mas distinto de la app real de lo que es en
realidad: la app real ya usa este mismo patron async, esta v2 al principio
no (fue la primera pagina copiada, antes de establecer esta costumbre)."""
import json
import os
import uuid
from datetime import datetime, timedelta
from io import BytesIO

import requests
from fastapi import APIRouter, Depends, Form
from fastapi.requests import Request
from fastapi.responses import JSONResponse, StreamingResponse

from .. import auth, db, logging_setup, mantis_client
from ..templating import templates

router = APIRouter(prefix="/mantis", tags=["mantis"])


def _empresas_mantis() -> list:
    return [e for e in db.load_empresas() if e["mantis_project_id"]]


def usuario_map() -> dict:
    """{usuario/login de Mantis -> user_id local}, guardado por Configuracion
    -> Mantis (ajustes_app, clave 'mantis.usuario_map', ver
    routers/configuracion.py) -- mismo patron que
    mantis_csv_import.mantis_csv_estado_map. Fase 1 del multiusuario (ver
    CLAUDE.md): mientras no haya una forma real de autenticar como cada
    persona en Mantis, esto es lo unico que permite asignar localmente una
    incidencia importada cuyo handler no sea 'yo'."""
    try:
        data = json.loads(db.get_ajuste("mantis.usuario_map", "{}") or "{}")
        return {str(k): int(v) for k, v in data.items() if k and v}
    except (ValueError, TypeError):
        return {}


def _mantis_base_url() -> str:
    return os.getenv("MANTIS_BASE_URL", "").strip().rstrip("/")


def _empresa_por_dominio(email: str, dominio_map: dict) -> str:
    """Sugiere una empresa comparando el dominio del email del reportero
    contra el dominio configurado en cada Empresa -- SOLO una sugerencia
    para precargar el selector manual, nunca se aplica en silencio (igual
    que core/views.py::_empresa_por_dominio: el dominio no es fiable al
    100%, gente de una empresa puede reportar incidencias de otra)."""
    dominio = (email or "").split("@")[-1].strip().lower()
    return dominio_map.get(dominio, "") if dominio else ""


@router.get("/")
def ver_mantis(request: Request, usuario: str = Depends(auth.requiere_login)):
    pool_id, _ = db.get_mantis_config()
    destino_id = db.get_ajuste("mantis.proyecto_destino", "")
    # el nombre del proyecto, no el id crudo -- antes esta v2 mostraba el id
    # tal cual en "Destino actual: ..." (bug visual encontrado al comparar
    # con el original, que siempre mostro active.nombre)
    destino_proyecto = db.get_proyecto(destino_id) if destino_id else None
    return templates.TemplateResponse(request, "mantis.html", {
        "usuario": usuario, "activo": "mantis",
        "mantis_destino_actual": destino_id, "mantis_destino_proyecto": destino_proyecto,
        "mantis_pool_project_id": pool_id,
        "empresas_mantis": _empresas_mantis(), "mantis_base_url": _mantis_base_url(),
    })


@router.get("/incidencias/json")
def incidencias_json(usuario: str = Depends(auth.requiere_login)):
    pool_id, empresa_project_ids = db.get_mantis_config()
    try:
        incidencias = mantis_client.fetch_pool_issues(pool_id, empresa_project_ids)
    except Exception as e:
        logging_setup.errores.exception(f"{usuario} | GET /mantis/incidencias/json -> fallo Mantis: {e}")
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)
    empresa_by_mantis_id = {int(e["mantis_project_id"]): e["empresa_id"] for e in _empresas_mantis()}
    dominio_map = {e["dominio_email"]: e["empresa_id"] for e in db.load_empresas() if e["dominio_email"]}
    for inc in incidencias:
        inc["es_pool"] = inc["proyecto_id"] == pool_id
        # solo tiene sentido sugerir empresa cuando el proyecto de Mantis en
        # el que esta ahora mismo no mapea ya a una -- si ya esta en una
        # empresa conocida esa ya se sabe seguro, sin adivinar nada
        inc["empresa_sugerida_id"] = (
            "" if inc["proyecto_id"] in empresa_by_mantis_id
            else _empresa_por_dominio(inc.get("reporta_email", ""), dominio_map)
        )
    return JSONResponse({"ok": True, "incidencias": incidencias})


@router.post("/importar/")
def importar_tareas(
    usuario: str = Depends(auth.requiere_login),
    issue_ids: str = Form("[]"), mover_a: str = Form(""), empresa_override: str = Form("{}"),
):
    proyecto_id = db.get_ajuste("mantis.proyecto_destino", "")
    if not proyecto_id:
        return JSONResponse({"ok": False, "error": "configura el proyecto destino en Configuración → Mantis"}, status_code=400)
    if not db.get_proyecto(proyecto_id):
        return JSONResponse({"ok": False, "error": "el proyecto destino configurado ya no existe"}, status_code=400)

    try:
        override_map = json.loads(empresa_override)
        if not isinstance(override_map, dict):
            override_map = {}
    except Exception:
        override_map = {}

    try:
        ids = {int(i) for i in json.loads(issue_ids)}
    except Exception:
        ids = set()
    if not ids:
        return JSONResponse({"ok": False, "error": "no se seleccionó ninguna incidencia"}, status_code=400)

    pool_id, empresa_project_ids = db.get_mantis_config()
    empresas_con_mapeo = _empresas_mantis()
    empresa_by_mantis_id = {int(e["mantis_project_id"]): e["empresa_id"] for e in empresas_con_mapeo}

    try:
        incidencias = mantis_client.fetch_pool_issues(pool_id, empresa_project_ids)
    except Exception as e:
        logging_setup.errores.exception(f"{usuario} | POST /mantis/importar/ -> fallo Mantis: {e}")
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)
    for inc in incidencias:
        inc["es_pool"] = inc["proyecto_id"] == pool_id
    seleccionadas = [i for i in incidencias if i["id"] in ids]

    movidas_ids = set()
    movidas_fallidas = []
    mover_a_id = None
    if mover_a.strip():
        try:
            mover_a_id = int(mover_a)
        except ValueError:
            mover_a_id = None
        if mover_a_id:
            for inc in seleccionadas:
                elegible = inc["es_pool"] and inc["asignada_a"] == "SIN ASIGNAR"
                if not elegible:
                    continue
                try:
                    mantis_client.move_issue(inc["id"], mover_a_id)
                    movidas_ids.add(inc["id"])
                except Exception as e:
                    logging_setup.errores.exception(f"{usuario} | move_issue #{inc['id']} -> {mover_a_id} fallo: {e}")
                    movidas_fallidas.append(f"#{inc['id']}: {e}")

    cat_by_nombre = {c["nombre"].strip().lower(): c["nombre"] for c in db.load_categorias_globales()}
    mantis_iniciativa_id = db.find_or_crear_iniciativa_mantis()
    ids_ya_importadas = db.tareas_mantis_ids_importadas()
    empresas_validas = {e["empresa_id"] for e in db.load_empresas()}
    mapa_usuarios = usuario_map()
    mi_user_id = db.get_user_id(usuario)

    creadas = 0
    ya_importadas = []
    for inc in seleccionadas:
        if inc["id"] in ids_ya_importadas:
            ya_importadas.append(inc["id"])
            continue
        categoria = cat_by_nombre.get((inc["categoria"] or "").strip().lower(), "")
        proyecto_id_final = mover_a_id if inc["id"] in movidas_ids else inc["proyecto_id"]
        empresa_id = empresa_by_mantis_id.get(proyecto_id_final, "")
        if not empresa_id:
            # sin mapeo de proyecto (sigue en el pool y no se movio): usa lo
            # que el usuario haya elegido a mano en el selector de la fila,
            # si eligio algo valido -- nunca se adivina en su lugar
            override = str(override_map.get(str(inc["id"]), ""))
            if override in empresas_validas:
                empresa_id = override
        pie = f"\n\n---\nMantis #{inc['id']} — {inc['url']}"
        # asignado_a: si el handler en Mantis soy yo, directo; si no, se
        # busca su usuario/login de Mantis en el mapeo de Configuracion ->
        # Mantis (por si ya se asocio a un auth_user local); sin ninguna de
        # las dos, se queda sin asignar en vez de adivinar.
        if inc.get("asignada_a_soy_yo"):
            asignado_a = mi_user_id
        else:
            asignado_a = mapa_usuarios.get(inc.get("asignada_a_username", ""))
        db.crear_tarea(
            str(uuid.uuid4())[:8], proyecto_id, categoria, inc["resumen"] or f"Incidencia #{inc['id']}",
            "Pendiente", inc["prioridad"], ((inc["descripcion"] or "") + pie).strip(), "",
            "", "", empresa_id=empresa_id, iniciativa_id=mantis_iniciativa_id, mantis_id=inc["id"],
            generada_por_ia=mantis_client.fue_creada_por_ia(inc["id"]),
            asignado_a=asignado_a,
        )
        creadas += 1

    return JSONResponse({
        "ok": True, "creadas": creadas,
        "movidas": len(movidas_ids), "movidas_fallidas": movidas_fallidas,
        "ya_importadas": ya_importadas,
    })


@router.post("/asignadas/sincronizar/")
def mantis_sincronizar_asignadas(usuario: str = Depends(auth.requiere_login)):
    """Boton "Sincronizar tareas asignadas a mí" de Configuración → Mantis:
    trae TODAS mis incidencias (cualquier estado, incluidas resueltas/
    cerradas) y crea una tarea local por cada una que todavía no esté
    importada -- para ponerse al día de golpe si se han ido acumulando sin
    traerlas una a una desde /mantis/ (que solo enseña las abiertas). Igual
    que core/views.py::mantis_sincronizar_asignadas. Las ya resueltas/
    cerradas en Mantis se crean YA en la columna "hecho" del proyecto
    destino, no en Pendiente. Único disparador, nada automático."""
    project_id = db.get_ajuste("mantis.proyecto_destino", "")
    if not project_id:
        return JSONResponse({"ok": False, "error": "configura el proyecto destino en Configuración → Mantis"}, status_code=400)
    if not db.get_proyecto(project_id):
        return JSONResponse({"ok": False, "error": "el proyecto destino configurado ya no existe"}, status_code=400)

    pool_id, empresa_project_ids = db.get_mantis_config()
    empresa_by_mantis_id = {int(e["mantis_project_id"]): e["empresa_id"] for e in _empresas_mantis()}
    try:
        incidencias = mantis_client.fetch_historico_mio([pool_id, *empresa_project_ids])
    except Exception as e:
        logging_setup.errores.exception(f"{usuario} | POST /mantis/asignadas/sincronizar/ -> fallo Mantis: {e}")
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)

    cat_by_nombre = {c["nombre"].strip().lower(): c["nombre"] for c in db.load_categorias_globales()}
    columnas = db.load_board_columns(project_id)
    done_col = next((c["nombre"] for c in columnas if c["es_done"]), "Hecha")
    mantis_iniciativa_id = db.find_or_crear_iniciativa_mantis()
    ids_ya_importadas = db.tareas_mantis_ids_importadas()

    creadas = 0
    for inc in incidencias:
        if inc["id"] in ids_ya_importadas:
            continue
        categoria = cat_by_nombre.get((inc["categoria"] or "").strip().lower(), "")
        empresa_id = empresa_by_mantis_id.get(inc["proyecto_id"], "")
        status = done_col if inc["estado_name"] in mantis_client.ESTADOS_CERRADOS else "Pendiente"
        pie = f"\n\n---\nMantis #{inc['id']} — {inc['url']}"
        # Este endpoint es "sincronizar MIS incidencias" (fetch_historico_mio
        # ya filtra por handler == mi_email) -- asignado_a es siempre quien
        # ejecuta la sincronizacion, sin condicion.
        db.crear_tarea(
            str(uuid.uuid4())[:8], project_id, categoria, inc["resumen"] or f"Incidencia #{inc['id']}",
            status, inc["prioridad"], ((inc["descripcion"] or "") + pie).strip(), "", "", "",
            empresa_id=empresa_id, iniciativa_id=mantis_iniciativa_id, mantis_id=inc["id"],
            generada_por_ia=mantis_client.fue_creada_por_ia(inc["id"]),
            asignado_a=db.get_user_id(usuario),
        )
        creadas += 1

    return JSONResponse({"ok": True, "creadas": creadas, "ya_importadas": len(incidencias) - creadas})


def _fetch_todos_los_proyectos_mantis() -> list[dict]:
    """Todos los proyectos que existen en Mantis (Administrar -> Proyectos),
    no solo los mapeados a una Empresa en Configuracion -> Empresas -- para
    el toggle 'Todas las empresas' de mantis_origen_desglose (pedido
    explicito del usuario: quiere ver tambien SAP/Intranet/PAS/IFES/FEAM/...,
    proyectos de Mantis que no son ninguna Empresa nuestra). No vive en
    core/mantis_client.py (nunca se toca Django, ver CLAUDE.md) -- llamada
    nueva, propia de esta v2, con el mismo patron de auth/timeout que el
    resto del cliente de Mantis. Devuelve id+nombre deduplicados por id (el
    admin de Mantis puede listar un mismo proyecto varias veces, anidado bajo
    distintos padres, pero la API REST lo da una sola vez)."""
    base_url = mantis_client._base_url()
    headers = {"Authorization": mantis_client._token()}
    resp = requests.get(f"{base_url}/api/rest/projects/", headers=headers, timeout=15)
    resp.raise_for_status()
    vistos = set()
    proyectos = []
    for p in resp.json().get("projects", []):
        if p["id"] in vistos:
            continue
        vistos.add(p["id"])
        proyectos.append({"id": p["id"], "nombre": p["name"]})
    return proyectos


def _sincronizar_espejo_mantis() -> None:
    """Pone al dia db.mantis_espejo con lo que cambio en Mantis desde la
    ultima vez (ver mantis_client.fetch_actualizadas_desde). Se re-pide un
    margen de 5 min antes de la marca por si Mantis guardo algun cambio con
    un updated_at ligeramente anterior al ultimo visto -- repetir filas es
    inocuo (upsert). Espejo vacio = listado completo, una sola vez."""
    marca = db.get_mantis_espejo_marca()
    desde_iso = (datetime.fromisoformat(marca) - timedelta(minutes=5)).isoformat() if marca else None
    incidencias = mantis_client.fetch_actualizadas_desde(desde_iso)
    filas = []
    for inc in incidencias:
        creada = db.mantis_iso_a_naive(inc["creada"])
        if not creada or inc["proyecto_id"] is None:
            continue
        filas.append({
            "mantis_id": inc["id"], "proyecto_id": inc["proyecto_id"], "proyecto_nombre": inc["proyecto"],
            "creada": creada, "actualizada": db.mantis_iso_a_naive(inc["actualizada"]) or creada,
            "tecnico": inc.get("asignada_a") or "SIN ASIGNAR", "categoria": inc.get("categoria") or "",
        })
    db.upsert_mantis_espejo(filas)
    if incidencias:
        nueva = max(incidencias, key=lambda i: datetime.fromisoformat(i["actualizada"]))["actualizada"]
        if not marca or datetime.fromisoformat(nueva) > datetime.fromisoformat(marca):
            db.set_mantis_espejo_marca(nueva)


@router.get("/origen-desglose/")
def mantis_origen_desglose(request: Request, usuario: str = Depends(auth.requiere_login)):
    """Desglose de incidencias de Mantis por origen (IA/Correo/Manual), una
    fila por Empresa (con mapeo a Mantis) mas 'Total', reagrupado tambien
    por tecnico asignado y por categoria. Cubre TODO el equipo (cualquier
    tecnico), no solo lo mio, a diferencia del resto del SLA de Mantis.

    'todas' (bool, ?todas=1): por defecto solo se listan las empresas CON
    proyecto de Mantis asociado. Con todas=1 se cuentan TAMBIEN
    todos los demas PROYECTOS DE MANTIS que no son ninguna Empresa nuestra
    (SAP, Intranet, PAS, Fiora...), con el nombre del proyecto de Mantis. El
    pool sin triar (ajustes_app.mantis.pool_project_id) cuenta SIEMPRE, en
    los dos modos, como fila propia "Incidencias (sin triar)" (antes se
    excluia; pedido explicito del usuario, 2026-09-29: sin el, un dia normal
    salia con menos de la mitad de lo que entra). Cuando una incidencia se
    mueve del pool a una empresa, pasa a contar en la fila de esa empresa.

    De donde salen los datos: proyecto/tecnico/categoria/fecha de cada
    incidencia vienen de db.mantis_espejo, que se pone al dia con Mantis al
    principio de CADA calculo (_sincronizar_espejo_mantis: solo lo cambiado
    desde la ultima vez, normalmente 1 pagina). Asi siempre refleja el estado
    actual -- nuevas, movidas de proyecto, reasignadas -- sin volver a listar
    el historico entero (~1 min con todos los proyectos). Sustituye al
    antiguo "rango ya cubierto hoy" por proyecto, que servia datos viejos
    (no veia lo creado/movido despues del primer calculo del dia). La
    clasificacion IA/correo (la llamada cara, una por incidencia) si es
    inmutable y sigue cacheada en mantis_origen_cache; solo se clasifica lo
    que cae en el rango y aun no estaba clasificado."""
    qp = request.query_params
    desde = qp.get("desde", "").strip()
    hasta = qp.get("hasta", "").strip()
    if not desde or not hasta:
        return JSONResponse({"ok": False, "error": "faltan 'desde'/'hasta'"}, status_code=400)
    todas_las_empresas = qp.get("todas", "").strip() in ("1", "true", "on")

    empresas_con_mapeo = _empresas_mantis()
    if not empresas_con_mapeo:
        return JSONResponse(
            {"ok": False, "error": "no hay ninguna empresa con proyecto de Mantis asociado (Configuración → Empresas)"},
            status_code=400,
        )
    pool_id, _ = db.get_mantis_config()

    # orden de filas: el pool, las empresas mapeadas (por su orden) y, con
    # 'todas', el resto de proyectos de Mantis por nombre -- incluidos los
    # que no tienen ninguna incidencia en el rango (salen a 0)
    nombre_por_proyecto: dict[int, str] = {pool_id: "Incidencias (sin triar)"}
    for empresa in sorted(empresas_con_mapeo, key=lambda e: e["orden"]):
        nombre_por_proyecto[int(empresa["mantis_project_id"])] = empresa["nombre"]
    try:
        _sincronizar_espejo_mantis()
        if todas_las_empresas:
            for proyecto in sorted(_fetch_todos_los_proyectos_mantis(), key=lambda p: p["nombre"]):
                if proyecto["id"] not in nombre_por_proyecto:
                    nombre_por_proyecto[proyecto["id"]] = proyecto["nombre"]
    except Exception as e:
        logging_setup.errores.exception(f"{usuario} | GET /mantis/origen-desglose/ -> fallo Mantis: {e}")
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)

    incidencias = [
        r for r in db.get_mantis_espejo(desde, hasta)
        if todas_las_empresas or r["proyecto_id"] in nombre_por_proyecto
    ]
    for r in incidencias:  # proyecto de Mantis que no salio en la lista de proyectos
        nombre_por_proyecto.setdefault(r["proyecto_id"], r["proyecto_nombre"] or f"Proyecto {r['proyecto_id']}")

    cache = db.get_mantis_origen_cache([r["mantis_id"] for r in incidencias])
    nuevas_para_cache = []
    conteo_por_proyecto = {pid: {"ia": 0, "correo": 0, "manual": 0} for pid in nombre_por_proyecto}
    conteo_por_usuario: dict = {}
    conteo_por_categoria: dict = {}
    total = {"ia": 0, "correo": 0, "manual": 0}
    detalle: list[dict] = []
    for r in incidencias:
        origen = cache.get(r["mantis_id"])
        if origen is None:
            origen = mantis_client.clasificar_origen(r["mantis_id"])
            nuevas_para_cache.append({
                "mantis_id": r["mantis_id"], "proyecto_id": r["proyecto_id"], "creada": r["creada"],
                "es_ia": origen["es_ia"], "es_correo": origen["es_correo"], "categoria": r["categoria"],
                "tecnico": r["tecnico"], "proyecto_nombre": r["proyecto_nombre"],
            })
        clave = "ia" if origen["es_ia"] else ("correo" if origen["es_correo"] else "manual")
        categoria = r["categoria"] or "Sin categoría"
        conteo_por_proyecto[r["proyecto_id"]][clave] += 1
        conteo_por_usuario.setdefault(r["tecnico"], {"ia": 0, "correo": 0, "manual": 0})[clave] += 1
        conteo_por_categoria.setdefault(categoria, {"ia": 0, "correo": 0, "manual": 0})[clave] += 1
        total[clave] += 1
        # detalle SIN datos personales mas alla de nombre de tecnico/empresa
        # -- es lo que el frontend guarda como "cache" en memoria del navegador
        # para el cruce de filtros estilo Power BI (clic en una tabla filtra
        # las otras 2 SIN volver a pedir nada al servidor) y tambien para el
        # grafico de tendencia (fecha = solo el dia, "YYYY-MM-DD").
        detalle.append({"empresa": nombre_por_proyecto[r["proyecto_id"]], "tecnico": r["tecnico"],
                        "categoria": categoria, "origen": clave, "fecha": r["creada"][:10]})
    if nuevas_para_cache:
        db.upsert_mantis_origen_cache(nuevas_para_cache)

    filas = [
        {"empresa": nombre_por_proyecto[pid], **c, "total": sum(c.values())}
        for pid, c in conteo_por_proyecto.items()
    ]
    filas.append({"empresa": "Total", **total, "total": sum(total.values())})

    filas_usuario = [
        {"usuario": u, "ia": c["ia"], "correo": c["correo"], "manual": c["manual"], "total": sum(c.values())}
        for u, c in conteo_por_usuario.items()
    ]
    filas_usuario.sort(key=lambda f: f["total"], reverse=True)
    filas_usuario.append({"usuario": "Total", **total, "total": sum(total.values())})

    filas_categoria = [
        {"categoria": cat, "ia": c["ia"], "correo": c["correo"], "manual": c["manual"], "total": sum(c.values())}
        for cat, c in conteo_por_categoria.items()
    ]
    filas_categoria.sort(key=lambda f: f["total"], reverse=True)
    filas_categoria.append({"categoria": "Total", **total, "total": sum(total.values())})

    return JSONResponse({
        "ok": True, "desde": desde, "hasta": hasta,
        "filas": filas, "filas_usuario": filas_usuario, "filas_categoria": filas_categoria,
        # detalle = 1 fila por incidencia ya clasificada, la "cache" que usa el
        # frontend para el cruce de filtros sin volver a llamar al servidor.
        "detalle": detalle,
    })


@router.post("/origen-desglose/exportar/")
async def mantis_origen_desglose_exportar(request: Request, usuario: str = Depends(auth.requiere_login)):
    """Exporta a un .xlsx de 3 hojas (Empresa/Técnico/Categoría) EXACTAMENTE
    lo que el navegador tiene ya calculado y pintado en pantalla -- incluido
    el filtro cruzado, si el usuario tiene uno activo (ver
    _renderMantisOrigenTablas()/_aggDimension() en templates/sla.html, que
    son quienes mandan estas 3 listas ya agregadas en el body). No repite
    ningun calculo ni llama a Mantis ni a la BD -- solo maqueta lo recibido
    como libro de Excel, mismo patron/estilo que
    routers/tablero.py::plantilla_importar_excel."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    body = await request.json()
    hojas = [
        ("Empresa", body.get("empresa") or []),
        ("Técnico", body.get("tecnico") or []),
        ("Categoría", body.get("categoria") or []),
    ]
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre_hoja, filas_hoja in hojas:
        ws = wb.create_sheet(nombre_hoja)
        ws.append([nombre_hoja, "IA", "Correo", "Manual", "Total"])
        for col in range(1, 6):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for f in filas_hoja:
            es_total = f.get("clave") == "Total"
            ws.append([f.get("clave", ""), f.get("ia", 0), f.get("correo", 0), f.get("manual", 0), f.get("total", 0)])
            if es_total:
                for col in range(1, 6):
                    ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col, width in zip("ABCDE", [30, 10, 10, 10, 10]):
            ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="desglose_origen_mantis.xlsx"'},
    )
