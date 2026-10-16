# -*- coding: utf-8 -*-
"""Indexa una carpeta (local o de red) en el catalogo de Nexus
(nexus_carpetas / nexus_archivos): nombre, ruta, fechas, tamaño y el TEXTO
de los archivos que se pueden leer (txt/md/csv/codigo, docx, pdf, xlsx), que
es lo que busca el FULLTEXT. Los archivos NO se copian: Nexus los enlaza en
su sitio, asi que la ruta tiene que ser legible tambien por la app (en
Docker, montar la carpeta en el contenedor con la misma ruta).

  python scripts/nexus_indexar.py "\\\\servidor\\documentacion"
  python scripts/nexus_indexar.py /srv/docs --tecnico-auto --limpiar
  python scripts/nexus_indexar.py docs --test

Re-ejecutable: lo ya indexado se actualiza (por ruta), no se duplica.
--limpiar quita del catalogo lo que ya no existe bajo esa carpeta.
--tecnico-auto marca el codigo/config/logs como rol 'tecnico': no sale en la
busqueda normal, solo en la "Busqueda dev" (ver nexus/search.py).
"""
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

from _comun import conectar

TEXTO = {".txt", ".md", ".csv", ".log", ".ini", ".cfg", ".conf", ".yml", ".yaml", ".json", ".xml", ".html", ".htm",
         ".sql", ".py", ".js", ".ts", ".ps1", ".sh", ".bat", ".php", ".css"}
TECNICO = {".sql", ".py", ".js", ".ts", ".ps1", ".sh", ".bat", ".php", ".css", ".json", ".xml", ".yml", ".yaml",
           ".ini", ".cfg", ".conf", ".log"}
TIPOS = {".pdf": "PDF", ".docx": "Documento", ".doc": "Documento", ".odt": "Documento", ".md": "Documento",
         ".txt": "Texto", ".xlsx": "Hoja de cálculo", ".xls": "Hoja de cálculo", ".csv": "Hoja de cálculo",
         ".pptx": "Presentación", ".ppt": "Presentación", ".drawio": "Diagrama", ".vsdx": "Diagrama",
         ".png": "Imagen", ".jpg": "Imagen", ".jpeg": "Imagen", ".msg": "Correo", ".eml": "Correo"}
IGNORAR = {"thumbs.db", "desktop.ini", ".ds_store"}
MAX_TEXTO = 2_000_000  # caracteres guardados por archivo


def extraer_texto(ruta: Path, ext: str, max_mb: float) -> str | None:
    try:
        if ruta.stat().st_size > max_mb * 1024 * 1024:
            return None
        if ext in TEXTO:
            return ruta.read_text(encoding="utf-8", errors="replace")[:MAX_TEXTO]
        if ext == ".docx":
            import docx
            d = docx.Document(str(ruta))
            partes = [p.text for p in d.paragraphs]
            for tabla in d.tables:
                partes += [" | ".join(c.text for c in fila.cells) for fila in tabla.rows]
            return "\n".join(partes)[:MAX_TEXTO]
        if ext == ".pdf":
            import pymupdf
            with pymupdf.open(str(ruta)) as pdf:
                return "\n".join(pagina.get_text() for pagina in pdf)[:MAX_TEXTO]
        if ext == ".xlsx":
            import openpyxl
            libro = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
            filas = []
            for hoja in libro.worksheets:
                for fila in hoja.iter_rows(values_only=True):
                    filas.append(" ".join(str(v) for v in fila if v is not None))
                    if sum(map(len, filas)) > MAX_TEXTO:
                        break
            return "\n".join(filas)[:MAX_TEXTO]
    except Exception as e:  # un archivo raro no para la indexacion
        print(f"  ! sin texto: {ruta} ({e.__class__.__name__})")
    return None


def _fecha(ts: float) -> str:
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("carpeta")
    ap.add_argument("--test", action="store_true", help="indexar en DB_NAME_TEST")
    ap.add_argument("--tecnico-auto", action="store_true", help="codigo/config/logs como rol 'tecnico'")
    ap.add_argument("--limpiar", action="store_true", help="quitar lo que ya no existe bajo la carpeta")
    ap.add_argument("--max-mb", type=float, default=30, help="no extraer texto de archivos mayores (MB)")
    args = ap.parse_args()

    raiz = Path(args.carpeta).resolve()
    if not raiz.is_dir():
        sys.exit(f"No existe la carpeta {raiz}")
    base = raiz.parent

    def rel(p: Path) -> str:
        return p.relative_to(base).as_posix()

    conn = conectar(os.getenv("DB_NAME_TEST" if args.test else "DB_NAME"))
    vistos_carpetas, vistos_archivos, n_texto = set(), set(), 0
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT GREATEST(COALESCE((SELECT MAX(id) FROM nexus_carpetas), 0), "
                        "COALESCE((SELECT MAX(id) FROM nexus_archivos), 0))")
            siguiente = [cur.fetchone()[0] + 1]

            def id_de(tabla: str, ruta_relativa: str) -> int:
                cur.execute(f"SELECT id FROM {tabla} WHERE ruta_relativa=%s", (ruta_relativa,))
                fila = cur.fetchone()
                if fila:
                    return fila[0]
                siguiente[0] += 1
                return siguiente[0] - 1

            ids_carpeta = {}
            for actual, subcarpetas, archivos in os.walk(raiz):
                subcarpetas[:] = sorted(d for d in subcarpetas if not d.startswith("."))
                carpeta = Path(actual)
                r = rel(carpeta)
                cid = id_de("nexus_carpetas", r)
                padre = ids_carpeta.get(carpeta.parent) if carpeta != raiz else None
                cur.execute(
                    "INSERT INTO nexus_carpetas (id, ruta_completa, ruta_relativa, nombre, carpeta_padre_id, nivel) "
                    "VALUES (%s,%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE ruta_completa=VALUES(ruta_completa), "
                    "carpeta_padre_id=VALUES(carpeta_padre_id), nivel=VALUES(nivel)",
                    (cid, str(carpeta), r, carpeta.name, padre, len(carpeta.relative_to(raiz).parts)),
                )
                ids_carpeta[carpeta] = cid
                vistos_carpetas.add(r)
                for nombre in sorted(archivos):
                    if nombre.lower() in IGNORAR or nombre.startswith("~$"):
                        continue
                    ruta = carpeta / nombre
                    ext = ruta.suffix.lower()
                    st = ruta.stat()
                    texto = extraer_texto(ruta, ext, args.max_mb)
                    n_texto += texto is not None
                    ra = rel(ruta)
                    cur.execute(
                        "INSERT INTO nexus_archivos (id, carpeta_id, ruta_completa, ruta_relativa, nombre, extension, "
                        "tamano_kb, fecha_modificacion, fecha_creacion, tipo_documento, contenido_plano, rol) "
                        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON DUPLICATE KEY UPDATE "
                        "carpeta_id=VALUES(carpeta_id), ruta_completa=VALUES(ruta_completa), tamano_kb=VALUES(tamano_kb), "
                        "fecha_modificacion=VALUES(fecha_modificacion), tipo_documento=VALUES(tipo_documento), "
                        "contenido_plano=VALUES(contenido_plano), rol=COALESCE(rol, VALUES(rol))",
                        (id_de("nexus_archivos", ra), cid, str(ruta), ra, nombre, ext or None,
                         round(st.st_size / 1024, 2), _fecha(st.st_mtime), _fecha(st.st_ctime),
                         TIPOS.get(ext, "Código" if ext in TECNICO else "Otro"), texto,
                         "tecnico" if args.tecnico_auto and ext in TECNICO else None),
                    )
                    vistos_archivos.add(ra)
            quitados = 0
            if args.limpiar:
                prefijo = rel(raiz)
                cur.execute("SELECT ruta_relativa FROM nexus_archivos WHERE ruta_relativa = %s OR ruta_relativa LIKE %s",
                            (prefijo, prefijo + "/%"))
                sobran = [r[0] for r in cur.fetchall() if r[0] not in vistos_archivos]
                for r in sobran:
                    cur.execute("DELETE FROM nexus_archivos WHERE ruta_relativa=%s", (r,))
                cur.execute("SELECT ruta_relativa FROM nexus_carpetas WHERE ruta_relativa = %s OR ruta_relativa LIKE %s",
                            (prefijo, prefijo + "/%"))
                sobran_c = [r[0] for r in cur.fetchall() if r[0] not in vistos_carpetas]
                for r in sorted(sobran_c, key=len, reverse=True):
                    cur.execute("DELETE FROM nexus_carpetas WHERE ruta_relativa=%s", (r,))
                quitados = len(sobran) + len(sobran_c)
        conn.commit()
    finally:
        conn.close()
    print(f"Indexado {raiz}: {len(vistos_carpetas)} carpeta(s), {len(vistos_archivos)} archivo(s), "
          f"{n_texto} con texto buscable" + (f", {quitados} quitado(s)" if args.limpiar else "") + ".")


if __name__ == "__main__":
    main()
